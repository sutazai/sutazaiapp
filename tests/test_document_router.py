"""
Test Document Router

This module contains tests for the document router.
"""

import os
import unittest
from pathlib import Path
from fastapi.testclient import TestClient
from fastapi import FastAPI
import pytest

import sys

sys.path.append("/opt/sutazaiapp")

from backend.routers.document_router import router


@pytest.mark.skip(reason="Document router endpoints not implemented yet.")
class TestDocumentRouter(unittest.TestCase):
    """Test cases for the Document Router."""

    def setUp(self):
        """Set up test environment."""
        self.app = FastAPI()
        self.app.include_router(router)
        self.client = TestClient(self.app)

        # Create test data directory if it doesn't exist
        self.test_data_dir = Path(__file__).parent / "data"
        self.test_data_dir.mkdir(parents=True, exist_ok=True)

        # Create test documents
        self._create_test_documents()

    def tearDown(self):
        """Clean up test environment."""
        # Remove test documents
        for file in self.test_data_dir.glob("*"):
            file.unlink()

    def _create_test_documents(self):
        """Create test documents for testing."""
        # Create a text file
        self.test_txt_path = self.test_data_dir / "test.txt"
        with open(self.test_txt_path, "w") as f:
            f.write("This is a test document.\nIt contains multiple lines.\n")

        # Create a PDF file
        try:
            import fpdf

            pdf = fpdf.FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            pdf.cell(200, 10, txt="This is a test PDF document.", ln=1)
            self.test_pdf_path = self.test_data_dir / "test.pdf"
            pdf.output(str(self.test_pdf_path))
        except ImportError:
            self.test_pdf_path = None

        # Create a DOCX file
        try:
            from docx import Document

            doc = Document()
            doc.add_paragraph("This is a test DOCX document.")
            self.test_docx_path = self.test_data_dir / "test.docx"
            doc.save(str(self.test_docx_path))
        except ImportError:
            self.test_docx_path = None

    def test_parse_document(self):
        """Test parsing a document."""
        if not os.path.exists(self.test_txt_path):
            self.skipTest("Test text file could not be created")

        with open(self.test_txt_path, "rb") as f:
            response = self.client.post(
                "/documents/parse",
                files={"file": ("test.txt", f, "text/plain")},
                params={"extract_text": True, "extract_metadata": True},
            )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check response structure
        self.assertIn("content", data)
        self.assertIn("metadata", data)

        # Check content
        self.assertIn("This is a test document", data["content"])
        self.assertIn("It contains multiple lines", data["content"])

        # Check metadata
        self.assertEqual(data["metadata"]["format"], "txt")
        self.assertGreater(data["metadata"]["size"], 0)

    def test_parse_nonexistent_file(self):
        """Test parsing a nonexistent file."""
        response = self.client.post(
            "/documents/parse",
            files={"file": ("nonexistent.txt", b"", "text/plain")},
            params={"extract_text": True, "extract_metadata": True},
        )

        self.assertEqual(response.status_code, 404)

    def test_parse_invalid_file(self):
        """Test parsing an invalid file."""
        # Create an invalid file
        invalid_file = self.test_data_dir / "invalid.bin"
        with open(invalid_file, "wb") as f:
            f.write(b"\x00\x01\x02\x03")

        try:
            with open(invalid_file, "rb") as f:
                response = self.client.post(
                    "/documents/parse",
                    files={"file": ("invalid.bin", f, "application/octet-stream")},
                    params={"extract_text": True, "extract_metadata": True},
                )

            self.assertEqual(response.status_code, 400)

        finally:
            if os.path.exists(invalid_file):
                os.remove(invalid_file)

    def test_get_supported_formats(self):
        """Test getting supported formats."""
        response = self.client.get("/documents/supported-formats")

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check response structure
        self.assertIn("formats", data)
        self.assertIn("max_file_size", data)

        # Check supported formats
        self.assertIn("txt", data["formats"])
        self.assertIn("pdf", data["formats"])
        self.assertIn("docx", data["formats"])

    def test_validate_document(self):
        """Test document validation."""
        if not os.path.exists(self.test_txt_path):
            self.skipTest("Test text file could not be created")

        with open(self.test_txt_path, "rb") as f:
            response = self.client.post(
                "/documents/validate", files={"file": ("test.txt", f, "text/plain")}
            )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check response structure
        self.assertIn("is_valid", data)
        self.assertIn("metadata", data)

        # Check validation result
        self.assertTrue(data["is_valid"])

    def test_extract_text(self):
        """Test text extraction."""
        if not os.path.exists(self.test_txt_path):
            self.skipTest("Test text file could not be created")

        with open(self.test_txt_path, "rb") as f:
            response = self.client.post(
                "/documents/extract-text", files={"file": ("test.txt", f, "text/plain")}
            )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check response structure
        self.assertIn("text", data)

        # Check extracted text
        self.assertIn("This is a test document", data["text"])
        self.assertIn("It contains multiple lines", data["text"])

    def test_extract_metadata(self):
        """Test metadata extraction."""
        if not os.path.exists(self.test_txt_path):
            self.skipTest("Test text file could not be created")

        with open(self.test_txt_path, "rb") as f:
            response = self.client.post(
                "/documents/extract-metadata",
                files={"file": ("test.txt", f, "text/plain")},
            )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check response structure
        self.assertIn("metadata", data)

        # Check metadata
        self.assertEqual(data["metadata"]["format"], "txt")
        self.assertGreater(data["metadata"]["size"], 0)
        self.assertIn("created", data["metadata"])
        self.assertIn("modified", data["metadata"])


if __name__ == "__main__":
    unittest.main()
